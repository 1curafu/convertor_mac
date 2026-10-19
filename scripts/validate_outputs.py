#!/usr/bin/env python3
"""Opens the files produced by the test suite with independent libraries,
so a broken XLSX/PPTX/DOCX/ODS/EPUB fails CI instead of failing in Office."""
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import docx
import openpyxl
import pptx

out = Path(sys.argv[1])
failures = []


def check(name, fn):
    try:
        detail = fn()
        print(f"✓ {name}: {detail}")
    except Exception as error:  # noqa: BLE001 - report every failure
        failures.append(name)
        print(f"✗ {name}: {error!r}")


def xlsx(path):
    book = openpyxl.load_workbook(path)
    sheet = book.worksheets[0]
    return f"{len(book.worksheets)} sheet(s), A1={sheet['A1'].value!r}, B3={sheet['B3'].value!r}"


def pptx_file(path):
    deck = pptx.Presentation(path)
    pictures = sum(1 for slide in deck.slides for shape in slide.shapes if shape.shape_type == 13)
    assert pictures == len(deck.slides) > 0, "every slide should hold one picture"
    return f"{len(deck.slides)} slide(s)"


def docx_file(path):
    document = docx.Document(path)
    text = " ".join(p.text for p in document.paragraphs)
    assert text.strip(), "no text"
    return text[:60]


def ods(path):
    with zipfile.ZipFile(path) as archive:
        first = archive.infolist()[0]
        assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED, "mimetype must be first and stored"
        ET.fromstring(archive.read("content.xml"))
        ET.fromstring(archive.read("META-INF/manifest.xml"))
    return "valid package"


def epub(path):
    with zipfile.ZipFile(path) as archive:
        first = archive.infolist()[0]
        assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED, "mimetype must be first and stored"
        assert archive.read("mimetype") == b"application/epub+zip"
        for name in archive.namelist():
            if name.endswith((".xhtml", ".opf", ".xml")):
                ET.fromstring(archive.read(name))  # EPUB content must be well-formed XML
    return f"{len(archive.namelist())} entries"


check("table.xlsx", lambda: xlsx(out / "table.xlsx"))
check("report.xlsx", lambda: xlsx(out / "report.xlsx"))
check("report.pptx", lambda: pptx_file(out / "report.pptx"))
check("report.docx", lambda: docx_file(out / "report.docx"))
check("notes.docx", lambda: docx_file(out / "notes.docx"))
check("table.ods", lambda: ods(out / "table.ods"))
check("notes.epub", lambda: epub(out / "notes.epub"))
check("report.epub", lambda: epub(out / "report.epub"))

if failures:
    sys.exit(f"{len(failures)} file(s) failed validation: {', '.join(failures)}")
